from __future__ import annotations

import csv
import io
import json
from dataclasses import dataclass, field
from pathlib import Path


@dataclass(frozen=True)
class CatalogItem:
    """One entry in a metadata catalog."""

    id: str
    metadata: dict[str, str] = field(default_factory=dict)


class Catalog:
    """A collection of metadata entries to curate from."""

    def __init__(self, items: list[CatalogItem]):
        self.items = items

    def __len__(self) -> int:
        return len(self.items)

    def __iter__(self):
        return iter(self.items)

    @classmethod
    def from_json(cls, path: str | Path) -> Catalog:
        return cls.from_json_text(Path(path).read_text(encoding="utf-8-sig"))

    @classmethod
    def from_json_text(cls, text: str) -> Catalog:
        data = json.loads(text.removeprefix("﻿"))
        items = []
        for index, entry in enumerate(data):
            entry = dict(entry)
            item_id = str(entry.pop("id")) if "id" in entry else str(index)
            items.append(CatalogItem(id=item_id, metadata=entry))
        return cls(items)

    @classmethod
    def from_csv(cls, path: str | Path) -> Catalog:
        return cls.from_csv_text(Path(path).read_text(encoding="utf-8-sig"))

    @classmethod
    def from_csv_text(cls, text: str) -> Catalog:
        items = []
        for index, row in enumerate(csv.DictReader(io.StringIO(text.removeprefix("﻿")))):
            row = dict(row)
            item_id = row.pop("id") if "id" in row else str(index)
            items.append(CatalogItem(id=item_id, metadata=row))
        return cls(items)

    @classmethod
    def from_excel(cls, path: str | Path) -> Catalog:
        return cls.from_excel_bytes(Path(path).read_bytes())

    @classmethod
    def from_excel_bytes(cls, data: bytes) -> Catalog:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        rows = workbook.active.iter_rows(values_only=True)
        headers = [str(h) for h in next(rows)]
        items = []
        for index, row in enumerate(rows):
            entry = {
                headers[i]: ("" if value is None else str(value))
                for i, value in enumerate(row)
                if i < len(headers)
            }
            item_id = entry.pop("id") if "id" in entry else str(index)
            items.append(CatalogItem(id=item_id, metadata=entry))
        return cls(items)
